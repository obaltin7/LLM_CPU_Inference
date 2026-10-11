"""Thesis and paper figures and tables from the Phase 2 campaign (sessions 01-03). Single entry point.

    python scripts/make_figures.py --lang tr      # Turkish, decimal comma  -> figures/tr/, tables/tr/
    python scripts/make_figures.py --lang en      # English, decimal point  -> figures/en/, tables/en/
    python scripts/make_figures.py --lang all

Data policy (decision of 2026-10-08): raw metrics (CPU J/token, tokens/s, gCO2eq/token) use all three
sessions; the NET energy (idle power subtracted) uses sessions 02-03 only, because the idle baseline of
session_01 was taken under 4.1 % background load. Error bars show the between-session spread
(+-1 standard deviation of the session means, n = 3) unless the caption says otherwise.

Figures: PDF (vector) and 300 dpi PNG. Colour-blind-safe palette; quantization is also encoded by
hatching / marker shape so the figures stay readable in black and white.
"""
import argparse
import json
import math
import os
import glob

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.lines import Line2D
from matplotlib.patches import Patch, Rectangle
from matplotlib.ticker import FuncFormatter, MultipleLocator
import numpy as np
import pandas as pd

import _paths  # noqa: F401  (makes src/ importable without installing the package)
from analyze_bandwidth import (PEAK_GFLOPS_FP32, bandwidth_models, bandwidth_threads, bytes_per_token, load_sizes,
                               memory_peak_from_metadata)
from analyze_campaign import GROUP, quantization_table, session_values, tost_table, variance_table
from analyze_night import (LADDER, load_night, quantization_ladder, thread_scaling_table, turbo_combined_ratio_table,
                           turbo_combined_table, turbo_ratio_table, turbo_table)
from llm_cpu_inference.models import MODELS, model_path
from llm_cpu_inference.config import PROJECT_ROOT

RESULTS_DIR = os.path.join(PROJECT_ROOT, "results")
NET_SESSIONS_EXCLUDED = {"session_01"}
MODEL_ORDER = ["1B-Q4_K_M", "1B-Q8_0", "3B-Q4_K_M", "3B-Q8_0", "8B-Q4_K_M", "8B-Q8_0",
               "Sparse8B-Q4_K_M", "Sparse8B-Q8_0"]
FAMILY_OF = {"1B": "1B", "3B": "3B", "8B": "8B", "Sparse8B": "Sparse 8B"}
PHYSICAL_CORES = 8

# Colour-blind-safe palette (validated); hatching / markers carry the same identity in B/W
QUANT_STYLE = {"Q4_K_M": {"color": "#2a78d6", "hatch": ""}, "Q8_0": {"color": "#eb6834", "hatch": "//"}}
METHOD_STYLE = {"quantization": {"color": "#2a78d6", "marker": "o"},
                "pruning+distillation": {"color": "#eb6834", "marker": "s"},
                "2:4 sparse pruning+distillation": {"color": "#1baf7a", "marker": "^"}}
SESSION_STYLE = {"session_01": {"color": "#2a78d6", "marker": "o"},
                 "session_02": {"color": "#eb6834", "marker": "s"},
                 "session_03": {"color": "#1baf7a", "marker": "^"}}
TURBO_STYLE = {"off": {"color": "#2a78d6", "hatch": ""}, "on": {"color": "#eb6834", "hatch": "//"}}
NIGHT_MARK = "†"  # dagger: measured in the night run (night_01), scaled through the anchor models
INK, MUTED, GRID = "#0b0b0b", "#52514e", "#e1e0d9"

STR = {
    "tr": {
        "decimal": ",", "csv_sep": ";",
        "quant": "nicemleme", "method": {"quantization": "nicemleme", "pruning+distillation": "budama + damıtma",
                                         "2:4 sparse pruning+distillation": "2:4 seyrek budama + damıtma"},
        "base": "temel (base) model", "pareto": "Pareto cephesi", "pareto_instruct": "Pareto cephesi (instruct modeller)",
        "J_per_token": "Enerji (J/token, CPU)", "J_per_token_net": "Net enerji (J/token, CPU)",
        "tps": "Üretim hızı (token/s)", "ppl": "Perplexity (WikiText-2)", "bpb": "Bit/bayt (WikiText-2)",
        "threads": "İş parçacığı sayısı", "phys_cores": "8 fiziksel çekirdek",
        "better_low": "← daha iyi", "better_high": "daha iyi →",
        "fig01_title_a": "Enerji verimliliği", "fig01_title_b": "Üretim hızı",
        "fig01_note": "Çubuk: üç oturumun ortalaması; hata çubuğu: oturum ortalamalarının ±1 standart sapması (n = 3 oturum, "
                      "oturum başına 10 tur × 3 istem). 8 iş parçacığı, 100 token/üretim, decode penceresi.",
        "fig02_note": "Yatay hata çubuğu: oturumlar arası ±1 SS (n = 3). Dikey hata çubuğu: perplexity için parça düzeyinde "
                      "%95 güven aralığı ({chunks} parça{full}).\nPareto cephesi yalnızca instruct modeller arasında; temel model "
                      "içi boş. " + NIGHT_MARK + " 8B-Q5_K_M ve 8B-Q6_K aynı gece (night_01, 10 tur) ölçüldü;\nJ/token "
                      "değerleri çapa modeller (8B-Q4_K_M, 8B-Q8_0) üzerinden kampanya ölçeğine normalize edildi.",
        "full_set": ", WikiText-2 testinin tamamı",
        "fig10_note_2s": "Çubuk: iki bağımsız oturumun (night_01, night_02; her biri 10 tur, tur = üç istem) ortalaması; hata "
                         "çubuğu oturumlar arası ±1 SS (n = 2). 8 iş parçacığı, 100 token/üretim; Turbo açık ve kapalı birimler "
                         "aynı tur içinde iç içe sıralıdır.\nTurbo kapalı: azami işlemci durumu %99 (ölçülen 3,50 GHz); açık: "
                         "%100 (ölçülen 4,80 GHz). Çubuk çiftinin üstündeki sayı açık/kapalı oranıdır (20 tur, eşleştirilmiş).",
        "fig10_title_a": "Enerji", "fig10_title_b": "Üretim hızı", "fig10_title_c": "Paket gücü",
        "power_W": "CPU paket gücü (W)", "turbo_on": "Turbo açık (%100)", "turbo_off": "Turbo kapalı (%99)",
        "fig10_note": "Çubuk: 10 turun ortalaması (tur = üç istem); hata çubuğu turlar arası ±1 SS. Tek oturum (night_01), "
                      "8 iş parçacığı, 100 token/üretim; Turbo açık ve kapalı birimler aynı tur içinde iç içe sıralıdır.\n"
                      "Turbo kapalı: azami işlemci durumu %99 (ölçülen 3,50 GHz); açık: %100 (ölçülen 4,80 GHz). "
                      "Çubuk çiftinin üstündeki sayı açık/kapalı oranıdır.",
        "fig11_title_a": "J/token, model içi en düşüğe göre (%)", "fig11_title_b": "Token/s, model içi en yükseğe göre (%)",
        "fig11_note": "Hücre: 5 turun ortalaması (tur = üç istem), tek oturum (night_01), 100 token/üretim; parantez içinde "
                      "mutlak değer (J/token, token/s). Kalın çerçeve: modelin en verimli (sol) ve en hızlı (sağ) ayarı.\n"
                      "Renk koyuluğu hücredeki sayıyla artar; sağ panelde en açık hücre en yavaş ayardır.",
        "fig03_note": "Nokta: üç oturumun ortalaması; hata çubuğu: oturum ortalamalarının ±1 SS (n = 3). Model 8B-Q4_K_M, "
                      "100 token/üretim. Kesikli çizgi: 8 fiziksel çekirdek (üstü Hyper-Threading).",
        "fig04_note": "Her nokta bir iş parçacığı sayısı (8B-Q4_K_M); hata çubukları oturumlar arası ±1 SS (n = 3). "
                      "Dolu noktalar Pareto-optimal.",
        "fig05_title_a": "Enerji deneyi (8 iş parçacığı)", "fig05_title_b": "İş parçacığı ölçekleme (8B-Q4_K_M)",
        "fig05_y": "Üç oturum ortalamasından sapma (%)",
        "fig05_note": "Her nokta bir oturumun ortalamasıdır (10 tur × 3 istem). Oturum 1 yeniden başlatma yapılmadan "
                      "ölçülmüştür (bkz. ek).",
        "fig06_note": "Çubuk: üç oturumun ortalaması; hata çubuğu: oturumlar arası ±1 SS (n = 3). Prefill 11–18 token'lık "
                      "istemler üzerinde ölçülmüştür (tek ileri geçiş, hesaplamaya bağlı); decode 100 token üretimidir. "
                      "Sağ panel: prefill penceresinin CPU enerjisi / istem token'ı ile decode J/token. 8 iş parçacığı.",
        "prefill": "Prefill (istem işleme)", "decode": "Decode (üretim)", "tps_axis": "Hız (token/s)",
        "prefill_J": "Prefill (J / istem token'ı)", "decode_J": "Decode (J / üretilen token)",
        "J_axis_prompt": "CPU enerjisi (J/token)",
        "session": "Oturum", "fig_names": {
            1: "fig01_energy_speed_per_model", 2: "fig02_energy_accuracy_pareto", 3: "fig03_thread_scaling",
            4: "fig04_thread_pareto", 5: "fig05_session_consistency", 6: "fig06_prefill_vs_decode",
            7: "fig07_bandwidth_per_model", 8: "fig08_bandwidth_per_thread", 9: "fig09_roofline",
            10: "fig10_turbo", 11: "fig11_thread_optimum"},
        "table_names": {1: "T1_main_results", 2: "T2_quantization_comparison", 3: "T3_sparse_vs_dense_equivalence",
                        4: "T4_thread_scaling", 5: "T5_session_variance_decomposition", 6: "T6_experimental_conditions",
                        7: "T6b_sessions", 8: "T7_memory_bandwidth", 9: "T8_prefill_vs_decode", 10: "T9_turbo",
                        11: "T10_thread_optimum"},
        "bw_axis": "Etkin bellek bant genişliği (GB/s)", "bw_peak": "Teorik tepe", "bw_est1": "tahmin 1 (tüm ağırlıklar)",
        "bw_est2": "tahmin 2 (gömme/KV düzeltmeli)",
        "fig07_note": "Çubuk: düzeltilmiş tahmin (ağırlıklar − satır bazlı okunan gömme tablosu + KV önbelleği) × decode token/s "
                      "(üç oturumun ortalaması, 8 iş parçacığı); hata çubuğu oturumlar arası ±1 SS. Kısa çizgi: tüm ağırlıkların "
                      "her token'da okunduğu varsayımı (tahmin 1). Yüzdeler teorik tepeye orandır.",
        "fig08_note": "8B-Q4_K_M, düzeltilmiş bayt/token tahmini × decode token/s; hata çubuğu oturumlar arası ±1 SS (n = 3). "
                      "Kesikli yatay çizgi: teorik tepe (2 kanal × 64 bit × 2133 MT/s).",
        "fig09_note": "Kaba roofline: aritmetik yoğunluk = 2 × parametre sayısı / düzeltilmiş bayt/token; elde edilen performans = "
                      "2 × parametre × token/s. Hesaplama çatısı AVX2 fp32 FMA üst sınırıdır (8 çekirdek × 3,5 GHz); nicemlenmiş "
                      "çekirdekler bu sınıra ulaşmaz, bu yüzden yalnızca bant genişliği çatısı nicel olarak anlamlıdır.",
        "roof_x": "Aritmetik yoğunluk (FLOP/bayt)", "roof_y": "Elde edilen performans (GFLOP/s)",
        "roof_bw": "Bant genişliği çatısı", "roof_compute": "Hesaplama çatısı (kaba üst sınır)",
        "equivalent": {"equivalent": "eşdeğer", "not equivalent": "eşdeğer değil", "inconclusive": "belirsiz"},
        "yes": "evet", "no": "hayır",
    },
    "en": {
        "decimal": ".", "csv_sep": ",",
        "quant": "quantization", "method": {"quantization": "quantization", "pruning+distillation": "pruning + distillation",
                                            "2:4 sparse pruning+distillation": "2:4 sparse pruning + distillation"},
        "base": "base model", "pareto": "Pareto frontier", "pareto_instruct": "Pareto frontier (instruct models)",
        "J_per_token": "Energy (J/token, CPU)", "J_per_token_net": "Net energy (J/token, CPU)",
        "tps": "Generation speed (tokens/s)", "ppl": "Perplexity (WikiText-2)", "bpb": "Bits per byte (WikiText-2)",
        "threads": "Number of threads", "phys_cores": "8 physical cores",
        "better_low": "← better", "better_high": "better →",
        "fig01_title_a": "Energy efficiency", "fig01_title_b": "Generation speed",
        "fig01_note": "Bar: mean of three sessions; error bar: ±1 SD of the session means (n = 3 sessions, "
                      "10 rounds × 3 prompts per session). 8 threads, 100 tokens per generation, decode window.",
        "fig02_note": "Horizontal error bar: between-session ±1 SD (n = 3). Vertical error bar: chunk-level 95 % CI of "
                      "the perplexity ({chunks} chunks{full}).\nThe Pareto frontier is drawn over instruct models only; the base "
                      "model is hollow. " + NIGHT_MARK + " 8B-Q5_K_M and 8B-Q6_K were measured in the same night (night_01, "
                      "10 rounds);\ntheir J/token is scaled to the campaign through the anchor models (8B-Q4_K_M, 8B-Q8_0).",
        "full_set": ", the whole WikiText-2 test set",
        "fig10_note_2s": "Bar: mean of two independent sessions (night_01, night_02; 10 rounds each, round = three prompts); "
                         "error bar: between-session ±1 SD (n = 2). 8 threads, 100 tokens per generation; Turbo-on and "
                         "Turbo-off units are interleaved within each round.\nTurbo off: maximum processor state 99 % (measured "
                         "3.50 GHz); on: 100 % (measured 4.80 GHz). The number above each pair is the on/off ratio (20 paired rounds).",
        "fig10_title_a": "Energy", "fig10_title_b": "Generation speed", "fig10_title_c": "Package power",
        "power_W": "CPU package power (W)", "turbo_on": "Turbo on (100 %)", "turbo_off": "Turbo off (99 %)",
        "fig10_note": "Bar: mean of 10 rounds (round = three prompts); error bar: between-round ±1 SD. One session (night_01), "
                      "8 threads, 100 tokens per generation; Turbo-on and Turbo-off units are interleaved within each round.\n"
                      "Turbo off: maximum processor state 99 % (measured 3.50 GHz); on: 100 % (measured 4.80 GHz). "
                      "The number above each pair is the on/off ratio.",
        "fig11_title_a": "J/token relative to the model's minimum (%)",
        "fig11_title_b": "Tokens/s relative to the model's maximum (%)",
        "fig11_note": "Cell: mean of 5 rounds (round = three prompts), one session (night_01), 100 tokens per generation; the "
                      "absolute value (J/token, tokens/s) is in brackets. Bold frame: the model's most frugal (left) and fastest "
                      "(right) setting.\nColour darkens with the number in the cell; in the right panel the lightest cell is the "
                      "slowest setting.",
        "fig03_note": "Point: mean of three sessions; error bar: ±1 SD of the session means (n = 3). Model 8B-Q4_K_M, "
                      "100 tokens per generation. Dashed line: 8 physical cores (Hyper-Threading beyond).",
        "fig04_note": "Each point is a thread count (8B-Q4_K_M); error bars are the between-session ±1 SD (n = 3). "
                      "Filled points are Pareto-optimal.",
        "fig05_title_a": "Energy experiment (8 threads)", "fig05_title_b": "Thread scaling (8B-Q4_K_M)",
        "fig05_y": "Deviation from the 3-session mean (%)",
        "fig05_note": "Each point is the mean of one session (10 rounds × 3 prompts). Session 1 was measured without a "
                      "reboot (see appendix).",
        "fig06_note": "Bar: mean of three sessions; error bar: between-session ±1 SD (n = 3). Prefill is measured on "
                      "11–18-token prompts (one forward pass, compute-bound); decode is the 100-token generation. "
                      "Right panel: CPU energy of the prefill window per prompt token next to the decode J/token. 8 threads.",
        "prefill": "Prefill (prompt processing)", "decode": "Decode (generation)", "tps_axis": "Speed (tokens/s)",
        "prefill_J": "Prefill (J / prompt token)", "decode_J": "Decode (J / generated token)",
        "J_axis_prompt": "CPU energy (J/token)",
        "session": "Session", "fig_names": {
            1: "fig01_energy_speed_per_model", 2: "fig02_energy_accuracy_pareto", 3: "fig03_thread_scaling",
            4: "fig04_thread_pareto", 5: "fig05_session_consistency", 6: "fig06_prefill_vs_decode",
            7: "fig07_bandwidth_per_model", 8: "fig08_bandwidth_per_thread", 9: "fig09_roofline",
            10: "fig10_turbo", 11: "fig11_thread_optimum"},
        "table_names": {1: "T1_main_results", 2: "T2_quantization_comparison", 3: "T3_sparse_vs_dense_equivalence",
                        4: "T4_thread_scaling", 5: "T5_session_variance_decomposition", 6: "T6_experimental_conditions",
                        7: "T6b_sessions", 8: "T7_memory_bandwidth", 9: "T8_prefill_vs_decode", 10: "T9_turbo",
                        11: "T10_thread_optimum"},
        "bw_axis": "Effective memory bandwidth (GB/s)", "bw_peak": "Theoretical peak", "bw_est1": "estimate 1 (all weights)",
        "bw_est2": "estimate 2 (embedding/KV corrected)",
        "fig07_note": "Bar: corrected estimate (weights − row-wise embedding table + KV cache) × decode tokens/s (mean of three "
                      "sessions, 8 threads); error bar: between-session ±1 SD. Tick: all weights read every token (estimate 1). "
                      "Percentages are the share of the theoretical peak.",
        "fig08_note": "8B-Q4_K_M, corrected bytes/token × decode tokens/s; error bar: between-session ±1 SD (n = 3). "
                      "Dashed line: theoretical peak (2 channels × 64 bit × 2133 MT/s).",
        "fig09_note": "Coarse roofline: arithmetic intensity = 2 × parameters / corrected bytes per token; attained performance = "
                      "2 × parameters × tokens/s. The compute roof is the AVX2 fp32 FMA upper bound (8 cores × 3.5 GHz); quantized "
                      "kernels do not reach it, so only the bandwidth roof is quantitatively meaningful.",
        "roof_x": "Arithmetic intensity (FLOP/byte)", "roof_y": "Attained performance (GFLOP/s)",
        "roof_bw": "Bandwidth roof", "roof_compute": "Compute roof (coarse upper bound)",
        "equivalent": {"equivalent": "equivalent", "not equivalent": "not equivalent", "inconclusive": "inconclusive"},
        "yes": "yes", "no": "no",
    },
}


# ---------------------------------------------------------------- formatting helpers
class Fmt:
    def __init__(self, lang):
        self.lang = lang
        self.s = STR[lang]

    def num(self, x, nd=2):
        if x is None or (isinstance(x, float) and math.isnan(x)):
            return "—"
        text = f"{x:,.{nd}f}"
        if self.s["decimal"] == ",":
            text = text.replace(",", "\x00").replace(".", ",").replace("\x00", ".")
        return text

    def pm(self, mean, sd, nd=2):
        return f"{self.num(mean, nd)} ± {self.num(sd, nd)}"

    def ci(self, low, high, nd=2):
        return f"[{self.num(low, nd)}; {self.num(high, nd)}]"

    def axis(self, ax, which="y", nd=1):
        formatter = FuncFormatter(lambda v, _: self.num(v, nd))
        (ax.yaxis if which == "y" else ax.xaxis).set_major_formatter(formatter)

    def method(self, m):
        return self.s["method"].get(m, m)


def style_axes(ax):
    ax.grid(color=GRID, linewidth=0.8, zorder=0)
    ax.tick_params(colors=MUTED, length=0)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color(MUTED)


def save(fig, out_dir, name):
    os.makedirs(out_dir, exist_ok=True)
    paths = []
    for ext in ("pdf", "png"):
        path = os.path.join(out_dir, f"{name}.{ext}")
        fig.savefig(path, dpi=300, facecolor="white")
        paths.append(path)
    plt.close(fig)
    return paths


def pareto_mask(speed, energy):
    """True where no other point is at least as fast AND at least as frugal (and better in one)."""
    s, e = list(speed), list(energy)
    return [not any(s[j] >= s[i] and e[j] <= e[i] and (s[j] > s[i] or e[j] < e[i]) for j in range(len(s)) if j != i)
            for i in range(len(s))]


# ---------------------------------------------------------------- data
class Data:
    def __init__(self, results_dir):
        self.energy = pd.read_csv(os.path.join(results_dir, "energy_iterations.csv"))
        self.thread = pd.read_csv(os.path.join(results_dir, "thread_scaling_iterations.csv"))
        self.cognitive = pd.read_csv(os.path.join(results_dir, "cognitive_results.csv")).set_index("model_key")
        # Prefill energy per prompt token (the prefill window of each measurement divided by its prompt length)
        self.energy["prefill_J_per_prompt_token"] = self.energy["prefill_cpu_J"] / self.energy["prompt_tokens"]
        self.metadata = {}
        for path in glob.glob(os.path.join(results_dir, "metadata", "*.json")):
            meta = json.load(open(path, encoding="utf-8"))
            if not meta.get("smoke"):
                self.metadata[meta["session_id"]] = meta
        self.sessions = sorted(self.energy["session_id"].unique())
        self.n_sessions = len(self.sessions)
        # Memory-bandwidth analysis (analyze_bandwidth.py); tensor sizes come from the cached CSV
        sizes = load_sizes(os.path.join(results_dir, "analysis"))
        self.peak_gbs = float(memory_peak_from_metadata(results_dir).loc[0, "theoretical_peak_gbs"])
        self.bw_models = bandwidth_models(sizes, self.energy, self.peak_gbs).set_index("model_key")
        self.bw_threads = bandwidth_threads(sizes, self.thread, self.peak_gbs).set_index("n_threads")
        self.bpt = bytes_per_token(sizes).set_index("model_key")
        # Night run (night_01, analyze_night.py): optional; the campaign figures are unchanged when it is absent
        self.cognitive_subset = self.cognitive  # 60-chunk subset of session_01 (appendix)
        self.night = load_night(results_dir)
        self.night2 = load_night(results_dir, "night_02")  # second run: B1 again + the full-set perplexity
        self.night_extra = []
        self.turbo_sessions = 0
        if self.night is not None:
            ne = self.night["energy"]
            quality, nll_dir = self.night["cognitive"], self.night["nll_dir"]
            if self.night2 is not None and len(self.night2["cognitive_full"]):
                quality, nll_dir = self.night2["cognitive_full"], self.night2["nll_dir"]  # 564 chunks
            self.night_ladder = quantization_ladder(ne, quality, sizes, self.peak_gbs, nll_dir).set_index("model_key")
            self.night_threads = thread_scaling_table(ne, sizes, self.peak_gbs)
            if self.night2 is not None:
                energies = {self.night["session_id"]: ne, self.night2["session_id"]: self.night2["energy"]}
                self.night_turbo = turbo_combined_table(energies)
                self.night_turbo_ratio = turbo_combined_ratio_table(energies).set_index("model_key")
                self.turbo_sessions = len(energies)
            else:
                self.night_turbo = turbo_table(ne)
                self.night_turbo_ratio = turbo_ratio_table(ne).set_index("model_key")
                self.turbo_sessions = 1
            if len(quality):
                self.cognitive = quality.set_index("model_key")  # full set when available, else 90 chunks
                self.cognitive_90 = self.night["cognitive"].set_index("model_key")
            self.ppl_chunks = int(self.cognitive["n_chunks"].iloc[0])
            # Night -> campaign scale through the two anchor models measured in both (mean of the two factors)
            camp_j = session_values(self.energy, "cpu_J_per_token").groupby("model_key")["cpu_J_per_token"].mean()
            camp_s = session_values(self.energy, "tokens_per_sec").groupby("model_key")["tokens_per_sec"].mean()
            anchors = [k for k in ("8B-Q4_K_M", "8B-Q8_0") if k in self.night_ladder.index and k in camp_j.index]
            self.anchor_J = float(np.mean([camp_j[k] / self.night_ladder.loc[k, "J_per_token"] for k in anchors]))
            self.anchor_speed = float(np.mean([camp_s[k] / self.night_ladder.loc[k, "tokens_per_sec"] for k in anchors]))
            self.night_extra = [k for k in LADDER if k in self.night_ladder.index and k not in camp_j.index]

    def night_energy(self, key):
        """Campaign-scale J/token (mean, sd) of a model measured only in the night run."""
        r = self.night_ladder.loc[key]
        return r["J_per_token"] * self.anchor_J, r["J_per_token_sd"] * self.anchor_J

    def night_speed(self, key):
        r = self.night_ladder.loc[key]
        return r["tokens_per_sec"] * self.anchor_speed, r["tokens_per_sec_sd"] * self.anchor_speed

    def stats(self, df, metric, index, net=False):
        """Mean and SD over the session means of `metric` for every configuration (index column)."""
        data = df[~df["session_id"].isin(NET_SESSIONS_EXCLUDED)] if net else df
        v = session_values(data, metric).pivot_table(index=index, columns="session_id", values=metric)
        return pd.DataFrame({"mean": v.mean(axis=1), "sd": v.std(axis=1, ddof=1), "n": v.count(axis=1)}).join(v)

    def energy_table(self):
        rows = {}
        for metric, net in (("cpu_J_per_token", False), ("cpu_J_per_token_net", True), ("tokens_per_sec", False),
                            ("gco2eq_per_token", False), ("prefill_tokens_per_sec", False), ("avg_cpu_power_W", False),
                            ("prefill_J_per_prompt_token", False)):
            rows[metric] = self.stats(self.energy, metric, "model_key", net)
        return rows

    def thread_table(self):
        rows = {}
        for metric, net in (("cpu_J_per_token", False), ("cpu_J_per_token_net", True), ("tokens_per_sec", False),
                            ("avg_cpu_power_W", False)):
            rows[metric] = self.stats(self.thread, metric, "n_threads", net)
        return rows


def model_label(key):
    family, quant = key.split("-", 1)
    return FAMILY_OF.get(family, family), quant


# ---------------------------------------------------------------- figures
def fig01_energy_speed(data, f, out_dir):
    e = data.energy_table()
    keys = [k for k in MODEL_ORDER if k in e["cpu_J_per_token"].index]
    x = np.arange(len(keys))
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.8))
    for ax, metric, title, ylabel, nd in ((axes[0], "cpu_J_per_token", f.s["fig01_title_a"], f.s["J_per_token"], 1),
                                         (axes[1], "tokens_per_sec", f.s["fig01_title_b"], f.s["tps"], 0)):
        t = e[metric].loc[keys]
        for i, key in enumerate(keys):
            quant = model_label(key)[1]
            st = QUANT_STYLE[quant]
            ax.bar(i, t.loc[key, "mean"], width=0.72, color=st["color"], hatch=st["hatch"], edgecolor="white",
                   linewidth=0.8, zorder=2)
        ax.errorbar(x, t["mean"], yerr=t["sd"], fmt="none", ecolor=INK, elinewidth=1.2, capsize=4, capthick=1.2, zorder=3)
        top = (t["mean"] + t["sd"]).max()
        for i, key in enumerate(keys):
            ax.text(i, t.loc[key, "mean"] + t.loc[key, "sd"] + top * 0.02, f.num(t.loc[key, "mean"], nd if nd else 1),
                    ha="center", va="bottom", fontsize=8.5, color=INK)
        ax.set_xticks(x)
        ax.set_xticklabels([model_label(k)[1] for k in keys], fontsize=8.5)
        # Family group labels under the quantization labels
        for fam in dict.fromkeys(model_label(k)[0] for k in keys):
            idx = [i for i, k in enumerate(keys) if model_label(k)[0] == fam]
            ax.annotate(fam, ((min(idx) + max(idx)) / 2, 0), xytext=(0, -26), textcoords="offset points",
                        ha="center", va="top", fontsize=10, color=INK, annotation_clip=False)
            if max(idx) < len(keys) - 1:
                ax.axvline(max(idx) + 0.5, color=GRID, linewidth=0.8, zorder=1)
        ax.set_title(title, color=INK, fontsize=12, pad=10)
        ax.set_ylabel(ylabel, color=MUTED)
        ax.set_ylim(0, top * 1.15)
        style_axes(ax)
        ax.grid(axis="x", visible=False)
        f.axis(ax, "y", nd)
    handles = [Patch(facecolor=QUANT_STYLE[q]["color"], hatch=QUANT_STYLE[q]["hatch"], edgecolor="white", label=q)
               for q in ("Q4_K_M", "Q8_0")]
    fig.legend(handles=handles, loc="lower center", ncol=2, frameon=False, bbox_to_anchor=(0.5, 0.075), fontsize=9,
               title=f.s["quant"].capitalize(), title_fontsize=9)
    fig.text(0.5, 0.005, f.s["fig01_note"], ha="center", va="bottom", color=MUTED, fontsize=7.5, wrap=True)
    fig.tight_layout(rect=(0, 0.16, 1, 1), w_pad=3)
    return save(fig, out_dir, f.s["fig_names"][1])


def fig02_tradeoff(data, f, out_dir):
    e = data.energy_table()["cpu_J_per_token"]
    cog = data.cognitive
    keys = [k for k in MODEL_ORDER if k in e.index and k in cog.index]
    extra = [k for k in data.night_extra if k in cog.index]  # 8B-Q5_K_M, 8B-Q6_K: night run, anchor-scaled
    if extra and "8B-Q4_K_M" in keys:
        i = keys.index("8B-Q4_K_M") + 1
        keys = keys[:i] + extra + keys[i:]
    xs_all = {k: (e.loc[k, "mean"] if k in e.index else data.night_energy(k)[0]) for k in keys}
    xerr = {k: (e.loc[k, "sd"] if k in e.index else data.night_energy(k)[1]) for k in keys}
    fig, axes = plt.subplots(1, 2, figsize=(11.5, 5.3))
    for ax, ycol, ylabel, lo, hi, nd in ((axes[0], "perplexity", f.s["ppl"], "ppl_ci95_low", "ppl_ci95_high", 1),
                                        (axes[1], "bits_per_byte", f.s["bpb"], None, None, 2)):
        xs = xs_all
        ys = {k: cog.loc[k, ycol] for k in keys}
        instruct = [k for k in keys if not MODELS[k]["base"]]
        optimal = pareto_mask([-xs[k] for k in instruct], [ys[k] for k in instruct])
        frontier = sorted([k for k, o in zip(instruct, optimal) if o], key=lambda k: xs[k])
        ax.plot([xs[k] for k in frontier], [ys[k] for k in frontier], color=MUTED, linewidth=1.5, zorder=1)
        for k in keys:
            st = METHOD_STYLE[MODELS[k]["method"]]
            base = MODELS[k]["base"]
            yerr = None
            if lo:
                yerr = [[ys[k] - cog.loc[k, lo]], [cog.loc[k, hi] - ys[k]]]
            ax.errorbar(xs[k], ys[k], xerr=xerr[k], yerr=yerr, fmt=st["marker"], markersize=9,
                        markerfacecolor="white" if base else st["color"], markeredgecolor=st["color"], markeredgewidth=2,
                        ecolor=st["color"], elinewidth=1.2, capsize=3, zorder=3)
            # The 8B ladder (Q4 < Q5 < Q6 < Q8) is dense in x: alternate label sides to avoid collisions
            offset = (8, -12 if base else 7)
            if k in extra:
                offset = (8, -14) if k.endswith("Q5_K_M") else (0, 23)
            elif k == "Sparse8B-Q4_K_M":
                offset = (-8, -12)  # to the left: the 8B ladder labels occupy the right-hand side
            ha = "center" if offset[0] == 0 else ("right" if offset[0] < 0 else "left")
            ax.annotate(k + (NIGHT_MARK if k in extra else ""), (xs[k], ys[k]), textcoords="offset points",
                        xytext=offset, ha=ha, fontsize=8.5, color=INK)
        ax.set_xlabel(f"{f.s['J_per_token']}  {f.s['better_low']}", color=MUTED)
        ax.set_ylabel(f"{ylabel}  {f.s['better_low']}", color=MUTED)
        ax.margins(x=0.18, y=0.18)
        style_axes(ax)
        ax.xaxis.set_major_locator(MultipleLocator(2))
        f.axis(ax, "x", 0)
        f.axis(ax, "y", nd)
    handles = [Line2D([], [], marker=st["marker"], linestyle="", markersize=9, color=st["color"], label=f.method(m))
               for m, st in METHOD_STYLE.items()]
    handles += [Line2D([], [], marker="o", linestyle="", markersize=9, markerfacecolor="white", markeredgecolor=MUTED,
                       markeredgewidth=2, label=f.s["base"]),
                Line2D([], [], color=MUTED, linewidth=1.5, label=f.s["pareto_instruct"])]
    fig.legend(handles=handles, loc="lower center", ncol=5, frameon=False, bbox_to_anchor=(0.5, 0.1), fontsize=8.5)
    chunks = getattr(data, "ppl_chunks", int(cog["n_chunks"].iloc[0]))
    note = f.s["fig02_note"].format(chunks=chunks, full=f.s["full_set"] if chunks >= 564 else "")
    fig.text(0.5, 0.005, note, ha="center", va="bottom", color=MUTED, fontsize=7.5)
    fig.tight_layout(rect=(0, 0.18, 1, 1), w_pad=3)
    return save(fig, out_dir, f.s["fig_names"][2])


def fig03_thread_scaling(data, f, out_dir):
    t = data.thread_table()
    threads = sorted(t["tokens_per_sec"].index)
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.6))
    for ax, metric, ylabel, nd, better in ((axes[0], "tokens_per_sec", f.s["tps"], 1, f.s["better_high"]),
                                          (axes[1], "cpu_J_per_token", f.s["J_per_token"], 1, f.s["better_low"])):
        s = t[metric].loc[threads]
        ax.errorbar(threads, s["mean"], yerr=s["sd"], fmt="o-", color="#2a78d6", markersize=7, linewidth=2,
                    ecolor=INK, elinewidth=1.2, capsize=4, zorder=3)
        ax.axvline(PHYSICAL_CORES, color=MUTED, linestyle="--", linewidth=1.2, zorder=1)
        ax.set_xscale("log", base=2)
        ax.set_xticks(threads)
        ax.set_xticklabels([str(n) for n in threads])
        ax.set_xlabel(f.s["threads"], color=MUTED)
        ax.set_ylabel(f"{ylabel}  {better}", color=MUTED)
        top = (s["mean"] + s["sd"]).max() * 1.15
        ax.set_ylim(0, top)
        # Label the core limit near the bottom, where no data point sits
        ax.text(PHYSICAL_CORES * 1.04, top * 0.04, f.s["phys_cores"], color=MUTED, fontsize=8.5, va="bottom", ha="left")
        style_axes(ax)
        f.axis(ax, "y", nd)
    fig.text(0.5, 0.01, f.s["fig03_note"], ha="center", va="bottom", color=MUTED, fontsize=7.5)
    fig.tight_layout(rect=(0, 0.06, 1, 1), w_pad=3)
    return save(fig, out_dir, f.s["fig_names"][3])


def fig04_thread_pareto(data, f, out_dir):
    t = data.thread_table()
    threads = sorted(t["tokens_per_sec"].index)
    speed, energy = t["tokens_per_sec"].loc[threads], t["cpu_J_per_token"].loc[threads]
    optimal = pareto_mask(speed["mean"], energy["mean"])
    fig, ax = plt.subplots(figsize=(7, 5))
    frontier = [n for n, o in zip(threads, optimal) if o]
    frontier.sort(key=lambda n: speed.loc[n, "mean"])
    ax.plot([speed.loc[n, "mean"] for n in frontier], [energy.loc[n, "mean"] for n in frontier], color="#2a78d6",
            linewidth=2, zorder=2, label=f.s["pareto"])
    for n, o in zip(threads, optimal):
        ax.errorbar(speed.loc[n, "mean"], energy.loc[n, "mean"], xerr=speed.loc[n, "sd"], yerr=energy.loc[n, "sd"],
                    fmt="o", markersize=9, markerfacecolor="#2a78d6" if o else "white", markeredgecolor="#2a78d6",
                    markeredgewidth=2, ecolor="#2a78d6", elinewidth=1.2, capsize=4, zorder=3)
        ax.annotate(f"{n} thread", (speed.loc[n, "mean"], energy.loc[n, "mean"]), textcoords="offset points",
                    xytext=(10, 7), fontsize=9, color=INK)
    ax.set_xlabel(f"{f.s['tps']}  {f.s['better_high']}", color=MUTED)
    ax.set_ylabel(f"{f.s['J_per_token']}  {f.s['better_low']}", color=MUTED)
    ax.margins(x=0.15, y=0.15)
    style_axes(ax)
    f.axis(ax, "x", 1)
    f.axis(ax, "y", 1)
    ax.legend(frameon=False, loc="lower left", fontsize=9)
    fig.text(0.5, 0.01, f.s["fig04_note"], ha="center", va="bottom", color=MUTED, fontsize=7.5)
    fig.tight_layout(rect=(0, 0.05, 1, 1))
    return save(fig, out_dir, f.s["fig_names"][4])


def fig05_session_consistency(data, f, out_dir):
    fig, axes = plt.subplots(1, 2, figsize=(11.5, 4.6), gridspec_kw={"width_ratios": [8, 7]})
    for ax, df, index, order, title in ((axes[0], data.energy, "model_key", MODEL_ORDER, f.s["fig05_title_a"]),
                                        (axes[1], data.thread, "n_threads", sorted(data.thread["n_threads"].unique()),
                                         f.s["fig05_title_b"])):
        v = session_values(df, "cpu_J_per_token").pivot_table(index=index, columns="session_id", values="cpu_J_per_token")
        v = v.loc[[o for o in order if o in v.index]]
        dev = 100 * (v.sub(v.mean(axis=1), axis=0)).div(v.mean(axis=1), axis=0)
        x = np.arange(len(dev))
        ax.axhline(0, color=MUTED, linewidth=1)
        for j, sid in enumerate(dev.columns):
            st = SESSION_STYLE[sid]
            ax.plot(x + (j - 1) * 0.12, dev[sid], linestyle="", marker=st["marker"], markersize=8, color=st["color"],
                    label=f"{f.s['session']} {int(sid[-2:])}", zorder=3)
        ax.set_xticks(x)
        ax.set_xticklabels([str(i) for i in dev.index], rotation=35 if index == "model_key" else 0, ha="right" if index == "model_key" else "center",
                           fontsize=8.5)
        ax.set_title(title, color=INK, fontsize=11, pad=8)
        ax.set_xlabel("" if index == "model_key" else f.s["threads"], color=MUTED)
        style_axes(ax)
        f.axis(ax, "y", 1)
    axes[0].set_ylabel(f.s["fig05_y"], color=MUTED)
    axes[1].legend(frameon=False, fontsize=9, loc="upper right")
    fig.text(0.5, 0.01, f.s["fig05_note"], ha="center", va="bottom", color=MUTED, fontsize=7.5)
    fig.tight_layout(rect=(0, 0.06, 1, 1), w_pad=3)
    return save(fig, out_dir, f.s["fig_names"][5])


def fig06_prefill_decode(data, f, out_dir):
    e = data.energy_table()
    keys = [k for k in MODEL_ORDER if k in e["tokens_per_sec"].index]
    x = np.arange(len(keys))
    fig, axes = plt.subplots(1, 2, figsize=(13, 5))
    panels = (
        (axes[0], f.s["tps_axis"], 0, 2, (("prefill_tokens_per_sec", f.s["prefill"]), ("tokens_per_sec", f.s["decode"]))),
        (axes[1], f.s["J_axis_prompt"], 2, 0.1,
         (("prefill_J_per_prompt_token", f.s["prefill_J"]), ("cpu_J_per_token", f.s["decode_J"]))),
    )
    for ax, ylabel, nd, pad, series in panels:
        top = 0
        for offset, (metric, label), color, hatch in zip((-0.19, 0.19), series, ("#1baf7a", "#2a78d6"), ("..", "")):
            t = e[metric].loc[keys]
            ax.bar(x + offset, t["mean"], width=0.36, color=color, hatch=hatch, edgecolor="white", linewidth=0.8,
                   label=label, zorder=2)
            ax.errorbar(x + offset, t["mean"], yerr=t["sd"], fmt="none", ecolor=INK, elinewidth=1.1, capsize=3, zorder=3)
            for i, key in enumerate(keys):
                ax.text(i + offset, t.loc[key, "mean"] + t.loc[key, "sd"] + pad, f.num(t.loc[key, "mean"], nd),
                        ha="center", va="bottom", fontsize=7, color=INK)
            top = max(top, (t["mean"] + t["sd"]).max())
        ax.set_xticks(x)
        ax.set_xticklabels(keys, rotation=30, ha="right", fontsize=8.5)
        ax.set_ylabel(ylabel, color=MUTED)
        ax.set_ylim(0, top * 1.14)
        style_axes(ax)
        ax.grid(axis="x", visible=False)
        f.axis(ax, "y", nd)
        ax.legend(frameon=False, fontsize=8.5, loc="upper right" if ax is axes[0] else "upper left")
    fig.text(0.5, 0.01, f.s["fig06_note"], ha="center", va="bottom", color=MUTED, fontsize=7.5, wrap=True)
    fig.tight_layout(rect=(0, 0.08, 1, 1), w_pad=3)
    return save(fig, out_dir, f.s["fig_names"][6])


def fig07_bandwidth_models(data, f, out_dir):
    bw = data.bw_models
    keys = [k for k in MODEL_ORDER if k in bw.index]
    x = np.arange(len(keys))
    fig, ax = plt.subplots(figsize=(10, 4.8))
    for i, key in enumerate(keys):
        st = QUANT_STYLE[model_label(key)[1]]
        ax.bar(i, bw.loc[key, "effective_gbs_est2"], width=0.66, color=st["color"], hatch=st["hatch"],
               edgecolor="white", linewidth=0.8, zorder=2)
        ax.plot([i - 0.33, i + 0.33], [bw.loc[key, "effective_gbs_est1"]] * 2, color=INK, linewidth=1.4, zorder=4)
        label_y = max(bw.loc[key, "effective_gbs_est2"] + bw.loc[key, "effective_gbs_est2_sd"],
                      bw.loc[key, "effective_gbs_est1"]) + 0.5
        ax.text(i, label_y, f"{f.num(bw.loc[key, 'utilisation_est2_pct'], 0)} %", ha="center", va="bottom",
                fontsize=8.5, color=INK)
    ax.errorbar(x, bw.loc[keys, "effective_gbs_est2"], yerr=bw.loc[keys, "effective_gbs_est2_sd"], fmt="none",
                ecolor=INK, elinewidth=1.2, capsize=4, zorder=3)
    ax.axhline(data.peak_gbs, color=MUTED, linestyle="--", linewidth=1.4, zorder=1)
    ax.text(len(keys) - 0.55, data.peak_gbs + 0.4, f"{f.s['bw_peak']} {f.num(data.peak_gbs, 1)} GB/s", ha="right",
            va="bottom", color=MUTED, fontsize=9)
    ax.set_xticks(x)
    ax.set_xticklabels(keys, rotation=30, ha="right", fontsize=8.5)
    ax.set_ylabel(f.s["bw_axis"], color=MUTED)
    ax.set_ylim(0, data.peak_gbs * 1.12)
    style_axes(ax)
    ax.grid(axis="x", visible=False)
    f.axis(ax, "y", 0)
    handles = [Patch(facecolor=QUANT_STYLE[q]["color"], hatch=QUANT_STYLE[q]["hatch"], edgecolor="white",
                     label=f"{q}: {f.s['bw_est2']}") for q in ("Q4_K_M", "Q8_0")]
    handles.append(Line2D([], [], color=INK, linewidth=1.4, label=f.s["bw_est1"]))
    fig.legend(handles=handles, frameon=False, fontsize=8.5, loc="lower center", ncol=3, bbox_to_anchor=(0.5, 0.075))
    fig.text(0.5, 0.005, f.s["fig07_note"], ha="center", va="bottom", color=MUTED, fontsize=7.5, wrap=True)
    fig.tight_layout(rect=(0, 0.15, 1, 1))
    return save(fig, out_dir, f.s["fig_names"][7])


def fig08_bandwidth_threads(data, f, out_dir):
    bw = data.bw_threads
    threads = sorted(bw.index)
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.errorbar(threads, bw.loc[threads, "effective_gbs_est2"], yerr=bw.loc[threads, "effective_gbs_est2_sd"],
                fmt="o-", color="#2a78d6", markersize=7, linewidth=2, ecolor=INK, elinewidth=1.2, capsize=4, zorder=3)
    for n in threads:
        ax.annotate(f"{f.num(bw.loc[n, 'utilisation_est2_pct'], 0)} %", (n, bw.loc[n, "effective_gbs_est2"]),
                    textcoords="offset points", xytext=(0, -16), ha="center", fontsize=8.5, color=INK)
    ax.axhline(data.peak_gbs, color=MUTED, linestyle="--", linewidth=1.4, zorder=1)
    ax.text(threads[-1], data.peak_gbs + 0.4, f"{f.s['bw_peak']} {f.num(data.peak_gbs, 1)} GB/s", ha="right",
            va="bottom", color=MUTED, fontsize=9)
    ax.axvline(PHYSICAL_CORES, color=MUTED, linestyle=":", linewidth=1.2, zorder=1)
    ax.set_xscale("log", base=2)
    ax.set_xticks(threads)
    ax.set_xticklabels([str(n) for n in threads])
    ax.set_xlabel(f.s["threads"], color=MUTED)
    ax.set_ylabel(f.s["bw_axis"], color=MUTED)
    ax.set_ylim(0, data.peak_gbs * 1.18)
    style_axes(ax)
    f.axis(ax, "y", 0)
    fig.text(0.5, 0.01, f.s["fig08_note"], ha="center", va="bottom", color=MUTED, fontsize=7.5, wrap=True)
    fig.tight_layout(rect=(0, 0.09, 1, 1))
    return save(fig, out_dir, f.s["fig_names"][8])


def fig09_roofline(data, f, out_dir):
    bw = data.bw_models
    keys = [k for k in MODEL_ORDER if k in bw.index]
    fig, ax = plt.subplots(figsize=(7.5, 5))
    xs = np.logspace(-1, 2.5, 200)
    roof = np.minimum(PEAK_GFLOPS_FP32, xs * data.peak_gbs)
    ax.plot(xs, roof, color=MUTED, linewidth=1.8, zorder=1)
    ridge = PEAK_GFLOPS_FP32 / data.peak_gbs
    ax.text(0.15, 0.15 * data.peak_gbs * 1.25, f.s["roof_bw"], color=MUTED, fontsize=8.5, rotation=36,
            rotation_mode="anchor", transform_rotates_text=True)
    ax.text(ridge * 2.5, PEAK_GFLOPS_FP32 * 1.12, f.s["roof_compute"], color=MUTED, fontsize=8.5, ha="center")
    for key in keys:
        st = METHOD_STYLE[MODELS[key]["method"]]
        base = MODELS[key]["base"]
        ax.plot(bw.loc[key, "arithmetic_intensity_flop_per_byte"], bw.loc[key, "attained_gflops"], linestyle="",
                marker=st["marker"], markersize=9, markerfacecolor="white" if base else st["color"],
                markeredgecolor=st["color"], markeredgewidth=2, zorder=3)
    # The eight models form two tight clusters (one per quantization); label the clusters, not the points
    for quant, dy in (("Q4_K_M", 1.6), ("Q8_0", 0.55)):
        members = [k for k in keys if k.endswith(quant)]
        cx = float(np.exp(np.log(bw.loc[members, "arithmetic_intensity_flop_per_byte"]).mean()))
        cy = float(bw.loc[members, "attained_gflops"].mean())
        ax.annotate(f"{quant} ({len(members)})", (cx, cy * dy), ha="center", va="bottom" if dy > 1 else "top",
                    fontsize=9, color=INK)
    ax.set_xscale("log")
    ax.set_yscale("log")
    ax.set_xlim(0.1, 300)
    ax.set_ylim(1, PEAK_GFLOPS_FP32 * 2.5)
    ax.set_xlabel(f.s["roof_x"], color=MUTED)
    ax.set_ylabel(f.s["roof_y"], color=MUTED)
    style_axes(ax)
    handles = [Line2D([], [], marker=st["marker"], linestyle="", markersize=9, color=st["color"], label=f.method(m))
               for m, st in METHOD_STYLE.items()]
    ax.legend(handles=handles, frameon=False, fontsize=8.5, loc="lower right")
    fig.text(0.5, 0.01, f.s["fig09_note"], ha="center", va="bottom", color=MUTED, fontsize=7.5, wrap=True)
    fig.tight_layout(rect=(0, 0.1, 1, 1))
    return save(fig, out_dir, f.s["fig_names"][9])


# ---------------------------------------------------------------- tables
def write_table(df, out_dir, name, f, title=None, note=None):
    """CSV (locale separators), aligned Markdown and a Word .docx table; `note` is printed under the table."""
    os.makedirs(out_dir, exist_ok=True)
    df = df.astype(str)
    df.to_csv(os.path.join(out_dir, f"{name}.csv"), index=False, sep=f.s["csv_sep"], encoding="utf-8-sig")

    widths = [max(len(str(c)), *(len(v) for v in df[c])) for c in df.columns]
    lines = ["| " + " | ".join(str(c).ljust(w) for c, w in zip(df.columns, widths)) + " |",
             "|" + "|".join("-" * (w + 2) for w in widths) + "|"]
    for _, row in df.iterrows():
        lines.append("| " + " | ".join(str(v).ljust(w) for v, w in zip(row, widths)) + " |")
    md = (f"**{title}**\n\n" if title else "") + "\n".join(lines) + "\n" + (f"\n*{note}*\n" if note else "")
    open(os.path.join(out_dir, f"{name}.md"), "w", encoding="utf-8").write(md)

    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:  # python-docx is optional; CSV and Markdown are always written
        return
    doc = Document()
    if title:
        doc.add_paragraph(title).runs[0].bold = True
    table = doc.add_table(rows=1, cols=len(df.columns))
    table.style = "Table Grid"
    for cell, col in zip(table.rows[0].cells, df.columns):
        cell.text = str(col)
        cell.paragraphs[0].runs[0].bold = True
    for _, row in df.iterrows():
        for cell, value in zip(table.add_row().cells, row):
            cell.text = str(value)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    if note:
        doc.add_paragraph(note).runs[0].font.size = Pt(8)
    doc.save(os.path.join(out_dir, f"{name}.docx"))


def table_T1(data, f):
    e = data.energy_table()
    cog = data.cognitive
    rows = []
    for k in MODEL_ORDER:
        if k not in e["cpu_J_per_token"].index:
            continue
        info = MODELS[k]
        # File size from the run metadata (recorded SHA-256/size), not from the local models/ directory
        size_gb = data.metadata[data.sessions[-1]]["models"][k]["size_bytes"] / 1e9
        c = cog.loc[k] if k in cog.index else None
        rows.append({
            ("Model" if f.lang == "en" else "Model"): k + (" (base)" if info["base"] else ""),
            ("Method" if f.lang == "en" else "Yöntem"): f.method(info["method"]),
            ("Params (B)" if f.lang == "en" else "Parametre (B)"): f.num(info["params_b"], 1),
            ("Quantization" if f.lang == "en" else "Nicemleme"): k.split("-", 1)[1],
            ("File (GB)" if f.lang == "en" else "Dosya (GB)"): f.num(size_gb, 2),
            ("Tokens/s" if f.lang == "en" else "Token/s"): f.pm(e["tokens_per_sec"].loc[k, "mean"], e["tokens_per_sec"].loc[k, "sd"], 2),
            ("J/token (raw)" if f.lang == "en" else "J/token (ham)"): f.pm(e["cpu_J_per_token"].loc[k, "mean"], e["cpu_J_per_token"].loc[k, "sd"], 3),
            ("J/token (net, S2–S3)" if f.lang == "en" else "J/token (net, O2–O3)"): f.pm(e["cpu_J_per_token_net"].loc[k, "mean"], e["cpu_J_per_token_net"].loc[k, "sd"], 3),
            "mgCO2eq/token": f.num(1000 * e["gco2eq_per_token"].loc[k, "mean"], 3),
            "Perplexity": f"{f.num(c['perplexity'], 3)} {f.ci(c['ppl_ci95_low'], c['ppl_ci95_high'], 2)}" if c is not None else "—",
            ("Bits/byte" if f.lang == "en" else "Bit/bayt"): f.num(c["bits_per_byte"], 4) if c is not None else "—",
        })
    return pd.DataFrame(rows)


def table_T2(data, f):
    subsets = {"all_sessions": data.sessions}
    q = quantization_table(data.energy, subsets)
    rows = []
    for _, r in q.iterrows():
        label = {"cpu_J_per_token": "J/token (ham)" if f.lang == "tr" else "J/token (raw)",
                 "cpu_J_per_token_net": "J/token (net)", "tokens_per_sec": "Token/s" if f.lang == "tr" else "Tokens/s"}[r["metric"]]
        if r["metric"] == "cpu_J_per_token_net":
            continue  # net ratios of session_01 are biased by its idle baseline; not reported
        rows.append({
            ("Family" if f.lang == "en" else "Aile"): r["comparison"].split(":")[0],
            ("Ratio" if f.lang == "en" else "Oran"): "Q8_0 / Q4_K_M" if "speed" not in r["comparison"] else "Q4_K_M / Q8_0",
            ("Metric" if f.lang == "en" else "Metrik"): label,
            ("Replicates" if f.lang == "en" else "Tekrarlar"): ", ".join(f.num(float(x), 3) for x in r["ratios_by_session"].split(";"))
                                                               + (" (3 sessions)" if f.lang == "en" else " (3 oturum)"),
            ("Mean" if f.lang == "en" else "Ortalama"): f.num(r["ratio_mean"], 3),
            ("95 % CI (Student-t)" if f.lang == "en" else "%95 GA (Student-t)"): f.ci(r["ratio_ci95_low"], r["ratio_ci95_high"], 3),
        })
    # Night run: the 8B ladder against the Q4_K_M anchor, paired per round (10 rounds, one session)
    if data.night is not None:
        nl = data.night_ladder
        rep = "10 rounds (night_01)" if f.lang == "en" else "10 tur (night_01)"
        for k in [k for k in LADDER if k in nl.index and k != "8B-Q4_K_M"]:
            q = k.split("-", 1)[1]
            r = nl.loc[k]
            rows.append({("Family" if f.lang == "en" else "Aile"): "8B " + NIGHT_MARK,
                         ("Ratio" if f.lang == "en" else "Oran"): f"{q} / Q4_K_M",
                         ("Metric" if f.lang == "en" else "Metrik"): "J/token (ham)" if f.lang == "tr" else "J/token (raw)",
                         ("Replicates" if f.lang == "en" else "Tekrarlar"): rep,
                         ("Mean" if f.lang == "en" else "Ortalama"): f.num(r["J_ratio_vs_anchor"], 3),
                         ("95 % CI (Student-t)" if f.lang == "en" else "%95 GA (Student-t)"): f.ci(r["J_ratio_ci95_low"], r["J_ratio_ci95_high"], 3)})
            rows.append({("Family" if f.lang == "en" else "Aile"): "8B " + NIGHT_MARK,
                         ("Ratio" if f.lang == "en" else "Oran"): f"Q4_K_M / {q}",
                         ("Metric" if f.lang == "en" else "Metrik"): "Token/s" if f.lang == "tr" else "Tokens/s",
                         ("Replicates" if f.lang == "en" else "Tekrarlar"): rep,
                         ("Mean" if f.lang == "en" else "Ortalama"): f.num(1 / r["speed_ratio_vs_anchor"], 3),
                         ("95 % CI (Student-t)" if f.lang == "en" else "%95 GA (Student-t)"): f.ci(1 / r["speed_ratio_ci95_high"], 1 / r["speed_ratio_ci95_low"], 3)})
    return pd.DataFrame(rows)


def table_T3(data, f):
    t = tost_table(data.energy, {"all_sessions": data.sessions})
    rows = []
    for _, r in t.iterrows():
        rows.append({
            ("Pair" if f.lang == "en" else "Çift"): r["pair"].replace(" vs ", " – "),
            ("Metric" if f.lang == "en" else "Metrik"): "J/token" if r["metric"] == "cpu_J_per_token" else ("Tokens/s" if f.lang == "en" else "Token/s"),
            ("Difference (%)" if f.lang == "en" else "Fark (%)"): f.num(r["mean_diff_pct"], 2),
            ("90 % CI (%)" if f.lang == "en" else "%90 GA (%)"): f.ci(r["ci90_low_pct"], r["ci90_high_pct"], 2),
            ("Margin" if f.lang == "en" else "Sınır"): f"±{f.num(r['margin_pct'], 0)} %",
            ("Verdict" if f.lang == "en" else "Karar"): f.s["equivalent"].get(r["verdict"], r["verdict"]),
        })
    return pd.DataFrame(rows)


def table_T4(data, f):
    t = data.thread_table()
    threads = sorted(t["tokens_per_sec"].index)
    optimal = pareto_mask(t["tokens_per_sec"].loc[threads, "mean"], t["cpu_J_per_token"].loc[threads, "mean"])
    rows = []
    for n, o in zip(threads, optimal):
        rows.append({
            ("Threads" if f.lang == "en" else "İş parçacığı"): n,
            ("Tokens/s" if f.lang == "en" else "Token/s"): f.pm(t["tokens_per_sec"].loc[n, "mean"], t["tokens_per_sec"].loc[n, "sd"], 3),
            ("J/token (raw)" if f.lang == "en" else "J/token (ham)"): f.pm(t["cpu_J_per_token"].loc[n, "mean"], t["cpu_J_per_token"].loc[n, "sd"], 3),
            ("J/token (net, S2–S3)" if f.lang == "en" else "J/token (net, O2–O3)"): f.pm(t["cpu_J_per_token_net"].loc[n, "mean"], t["cpu_J_per_token_net"].loc[n, "sd"], 3),
            ("Package power (W)" if f.lang == "en" else "Paket gücü (W)"): f.num(t["avg_cpu_power_W"].loc[n, "mean"], 1),
            ("Pareto-optimal" if f.lang == "en" else "Pareto-optimal"): f.s["yes"] if o else f.s["no"],
        })
    return pd.DataFrame(rows)


def table_T5(data, f):
    v = variance_table({"energy": data.energy, "thread_scaling": data.thread})
    v = v[v["subset"] == "all_sessions"]
    rows = []
    for _, r in v.iterrows():
        rows.append({
            ("Experiment" if f.lang == "en" else "Deney"): "energy" if r["experiment"] == "energy" else "thread scaling",
            ("Configuration" if f.lang == "en" else "Yapılandırma"): f"{r['model_key']} / {int(r['n_threads'])} thr",
            ("Metric" if f.lang == "en" else "Metrik"): "J/token" if r["metric"] == "cpu_J_per_token" else ("Tokens/s" if f.lang == "en" else "Token/s"),
            ("Mean" if f.lang == "en" else "Ortalama"): f.num(r["grand_mean"], 3),
            ("Within-session CV (%)" if f.lang == "en" else "Oturum içi CV (%)"): f.num(r["cv_within_pct"], 2),
            ("Between-session CV (%)" if f.lang == "en" else "Oturumlar arası CV (%)"): f.num(r["cv_between_pct"], 2),
            ("Between share (%)" if f.lang == "en" else "Oturumlar arası pay (%)"): f.num(r["between_share_pct"], 1),
        })
    return pd.DataFrame(rows)


def table_T6(data, f):
    meta = data.metadata[data.sessions[-1]]
    hw, sw, cfg = meta["hardware"], meta["software"], meta["config"]
    mem = hw["memory"]
    tr = f.lang == "tr"
    # Rated (XMP) speed from the part number, e.g. F4-3200C16-16GSXFB -> 3200; SPD reports the JEDEC speed
    import re
    part = mem["modules"][0]["part_number"]
    rated = re.search(r"-(\d{4})C", part)
    rated_text = (f"(XMP kapalı; anma {rated.group(1)} MHz*)" if tr else f"(XMP off; rated {rated.group(1)} MHz*)") if rated else ""
    ss, wl, idle = cfg["session_start"], cfg["warmup_load"], cfg["idle"]
    session_start_text = (f"açılıştan ≥{ss['min_uptime_s'] // 60} dk, sessizlik kapısı <%{f.num(ss['quiet_threshold_pct'], 0)} CPU "
                          f"({ss['quiet_window_s']} s), {idle['duration_s']} s boşta ölçüm, {wl['duration_s']} s ısınma yükü "
                          f"({wl['model']}, {wl['n_threads']})" if tr else
                          f"≥{ss['min_uptime_s'] // 60} min after boot, quiet gate <{f.num(ss['quiet_threshold_pct'], 0)} % CPU "
                          f"({ss['quiet_window_s']} s), {idle['duration_s']} s idle baseline, {wl['duration_s']} s warm-up load "
                          f"({wl['model']}, {wl['n_threads']})")
    sha_text = "; ".join(f"{k}: {v['sha256'][:12]}" for k, v in meta["models"].items())
    items = [
        ("CPU", f"{hw['cpu']['model']} ({hw['topology']['physical_cores']}C/{hw['topology']['logical_processors']}T, "
                f"L3 {hw['cpu']['l3_cache_kb'] // 1024} MB)"),
        ("RAM", f"{mem['total_gb']:.0f} GB DDR4, {len(mem['modules'])} × {mem['modules'][0]['capacity_gb']:.0f} GB "
                f"{part}, {mem['configured_speed_mhz']} MHz {rated_text}"),
        ("Anakart" if tr else "Mainboard", f"{hw['board']['manufacturer']} {hw['board']['product']}"),
        ("GPU", "; ".join(g["name"] for g in hw["gpu"]) + (" (çıkarımda kullanılmadı)" if tr else " (not used for inference)")),
        ("İşletim sistemi" if tr else "Operating system", f"{hw['os']['caption']} {hw['os']['version']}"),
        ("Güç planı" if tr else "Power plan", f"{hw['power']['name']} ({hw['power']['canonical']}), "
                                             + ("maks. işlemci durumu" if tr else "max processor state")
                                             + f" {hw['power']['max_processor_state_ac']} % → Turbo Boost "
                                             + ("kapalı" if tr else "off")),
        ("Python / llama-cpp-python / CodeCarbon", f"{sw['python']} / {sw['llama_cpp_python']} / {sw['codecarbon']}"),
        ("llama.cpp CPU özellikleri" if tr else "llama.cpp CPU features",
         ", ".join(p.strip().replace(" = 1", "") for p in sw["llama_cpp_system_info"].replace("CPU : ", "").split("|")
                   if p.strip().endswith("= 1"))),
        ("Enerji kaynağı" if tr else "Energy source", meta["energy_tracking"]["cpu_energy_scope"] + " — "
                                                      + meta["energy_tracking"]["cpu_energy_source"]),
        ("Karbon yoğunluğu" if tr else "Carbon intensity", f"{f.num(meta['carbon']['grid_intensity_gco2eq_per_kwh'], 1)} gCO2eq/kWh "
                                                           f"({meta['carbon']['country_iso_code']}, CodeCarbon)"),
        ("Bağlam / üretim" if tr else "Context / generation", f"n_ctx {cfg['inference']['n_ctx']}, {cfg['energy']['n_predict']} token, "
                                                               + ("açgözlü kod çözme, EOS bastırılmış" if tr else "greedy decoding, EOS suppressed")),
        ("Tur / istem / ısınma" if tr else "Rounds / prompts / warm-up", f"{cfg['energy']['iterations']} / {len(cfg['prompts'])} / {cfg['energy']['warmup']}"),
        ("Termal sigorta" if tr else "Thermal fuse", f"{cfg['thermal']['abort_c']} °C / {cfg['thermal']['abort_hold_s']} s; "
                                                   + ("soğuma hedefi" if tr else "cool-down target") + f" {cfg['thermal']['cooldown_target_c']} °C"),
        ("Oturum başlangıcı" if tr else "Session start", session_start_text),
        ("Model dosyaları (SHA-256 öneki)" if tr else "Model files (SHA-256 prefix)", sha_text),
    ]
    return pd.DataFrame(items, columns=["Koşul" if tr else "Condition", "Değer" if tr else "Value"])


T6_NOTE = {
    "tr": "(*) Anma hız modül parça numarasından türetilmiştir (XMP profili); SPD'nin bildirdiği JEDEC hızı metadata'daki "
          "configured/rated değerdir (2133 MT/s). Oturum başlangıcı değerleri configs/experiment.yaml'dan okunur.",
    "en": "(*) Rated speed is derived from the module part number (XMP profile); the SPD-reported JEDEC speed is the "
          "configured/rated value in the metadata (2133 MT/s). Session-start values are read from configs/experiment.yaml.",
}


def table_T6b(data, f):
    tr = f.lang == "tr"
    rows = []
    for sid in data.sessions:
        m = data.metadata[sid]
        idle = m.get("idle_baseline") or {}
        uptime = (m.get("session_start") or {}).get("uptime_at_start_s")
        start, end = m["timestamp_start"][:16].replace("T", " "), m["timestamp_end"][11:16]
        rows.append({
            ("Session" if not tr else "Oturum"): sid,
            ("Date / time" if not tr else "Tarih / saat"): f"{start} – {end}",
            ("Reboot before" if not tr else "Öncesinde yeniden başlatma"): (f.s["no"] if sid == "session_01" else f.s["yes"]),
            ("Uptime at start" if not tr else "Başlangıçta açık kalma"): (f"{uptime / 60:.0f} {'dk' if tr else 'min'}"
                                                                           if uptime is not None else ("~11,5 sa" if tr else "~11.5 h")),
            ("Idle CPU (W / %)" if not tr else "Boşta CPU (W / %)"): f"{f.num(idle.get('cpu_W'), 2)} / {f.num(idle.get('cpu_util_percent'), 1)}",
            ("Max package °C" if not tr else "En yüksek paket °C"): f.num((m.get("thermal") or {}).get("max_temperature_c"), 0),
            ("Room °C" if not tr else "Oda °C"): (m.get("operator_note") or "").replace("oda ", "").replace(" C", ""),
            ("Units" if not tr else "Birim"): len([u for u in m.get("executed_units", []) if u.get("unit_id") != "warmup_load"]),
        })
    return pd.DataFrame(rows)


def table_T7(data, f):
    bw = data.bw_models
    tr = f.lang == "tr"
    rows = []
    keys = [k for k in MODEL_ORDER if k in bw.index]
    if data.night_extra and "8B-Q4_K_M" in keys:
        i = keys.index("8B-Q4_K_M") + 1
        keys = keys[:i] + data.night_extra + keys[i:]
    for k in keys:
        if k in bw.index:
            r = bw.loc[k]
            speed, speed_sd = r["tokens_per_sec"], r["tokens_per_sec_sd"]
            label = k + (" (base)" if MODELS[k]["base"] else "")
        else:  # night-only model: campaign-scale speed through the anchors, bytes/token from the GGUF header
            b = data.bpt.loc[k]
            speed, speed_sd = data.night_speed(k)
            r = {"estimate_1_bytes": b["estimate_1_bytes"], "estimate_2_bytes": b["estimate_2_bytes"],
                 "effective_gbs_est1": b["estimate_1_bytes"] * speed / 1e9,
                 "effective_gbs_est2": b["estimate_2_bytes"] * speed / 1e9,
                 "effective_gbs_est2_sd": b["estimate_2_bytes"] * speed_sd / 1e9}
            r["utilisation_est2_pct"] = 100 * r["effective_gbs_est2"] / data.peak_gbs
            label = k + " " + NIGHT_MARK
        rows.append({
            "Model": label,
            ("Bayt/token, tahmin 1 (GB)" if tr else "Bytes/token, estimate 1 (GB)"): f.num(r["estimate_1_bytes"] / 1e9, 3),
            ("Bayt/token, tahmin 2 (GB)" if tr else "Bytes/token, estimate 2 (GB)"): f.num(r["estimate_2_bytes"] / 1e9, 3),
            ("Token/s" if tr else "Tokens/s"): f.pm(speed, speed_sd, 2),
            ("Etkin GB/s, tahmin 1" if tr else "Effective GB/s, estimate 1"): f.num(r["effective_gbs_est1"], 1),
            ("Etkin GB/s, tahmin 2" if tr else "Effective GB/s, estimate 2"): f.pm(r["effective_gbs_est2"], r["effective_gbs_est2_sd"], 1),
            ("Teorik tepeye oran (%)" if tr else "Share of theoretical peak (%)"): f.num(r["utilisation_est2_pct"], 1),
        })
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- night run (night_01): Turbo and thread optimum
def fig10_turbo(data, f, out_dir):
    t, r = data.night_turbo, data.night_turbo_ratio
    models = [k for k in MODEL_ORDER if k in set(t["model_key"])]
    x = np.arange(len(models))
    w = 0.38
    fig, axes = plt.subplots(1, 3, figsize=(12, 4.9))
    specs = ((axes[0], "J_per_token", f.s["fig10_title_a"], f.s["J_per_token"], "J_ratio_on_off", 1),
             (axes[1], "tokens_per_sec", f.s["fig10_title_b"], f.s["tps"], "speed_ratio_on_off", 1),
             (axes[2], "cpu_power_W", f.s["fig10_title_c"], f.s["power_W"], "power_ratio_on_off", 0))
    for ax, metric, title, ylabel, ratio_col, nd in specs:
        tops = []
        for j, state in enumerate(("off", "on")):
            st = TURBO_STYLE[state]
            sub = t[t["turbo"] == state].set_index("model_key").loc[models]
            pos = x + (j - 0.5) * w
            ax.bar(pos, sub[metric], width=w * 0.94, color=st["color"], hatch=st["hatch"], edgecolor="white",
                   linewidth=0.8, zorder=2)
            ax.errorbar(pos, sub[metric], yerr=sub[metric + "_sd"], fmt="none", ecolor=INK, elinewidth=1.1, capsize=3,
                        capthick=1.1, zorder=3)
            for p, v, sd in zip(pos, sub[metric], sub[metric + "_sd"]):
                ax.text(p, v + sd + 0.01 * sub[metric].max(), f.num(v, nd), ha="center", va="bottom", fontsize=7.5,
                        color=MUTED)
            tops.append(sub[metric] + sub[metric + "_sd"])
        top = pd.concat(tops, axis=1).max(axis=1)
        for i, k in enumerate(models):
            ax.text(i, top.iloc[i] + 0.08 * top.max(), "×" + f.num(r.loc[k, ratio_col], 3 if ratio_col == "speed_ratio_on_off" else 2),
                    ha="center", va="bottom", fontsize=9, color=INK, fontweight="bold")
        ax.set_xticks(x)
        ax.set_xticklabels(models, fontsize=8.5)
        ax.set_title(title, color=INK, fontsize=12, pad=10)
        ax.set_ylabel(ylabel, color=MUTED)
        ax.set_ylim(0, top.max() * 1.25)
        style_axes(ax)
        ax.grid(axis="x", visible=False)
        f.axis(ax, "y", nd)
    handles = [Patch(facecolor=TURBO_STYLE[s]["color"], hatch=TURBO_STYLE[s]["hatch"], edgecolor="white",
                     label=f.s["turbo_" + s]) for s in ("off", "on")]
    fig.legend(handles=handles, loc="lower center", ncol=2, frameon=False, bbox_to_anchor=(0.5, 0.1), fontsize=9)
    fig.text(0.5, 0.005, f.s["fig10_note_2s" if data.turbo_sessions >= 2 else "fig10_note"], ha="center", va="bottom",
             color=MUTED, fontsize=7.5)
    fig.tight_layout(rect=(0, 0.17, 1, 1), w_pad=3)
    return save(fig, out_dir, f.s["fig_names"][10])


def fig11_thread_optimum(data, f, out_dir):
    ts = data.night_threads
    models = [k for k in MODEL_ORDER if k in set(ts["model_key"])]
    threads = sorted(ts["n_threads"].unique())
    cmap = LinearSegmentedColormap.from_list("seq_blue", ["#ffffff", "#2a78d6"])
    fig, axes = plt.subplots(1, 2, figsize=(11, 5.4))
    panels = ((axes[0], "J_rel_to_min", "J_per_token", f.s["fig11_title_a"], 2, "min"),
              (axes[1], "speed_rel_to_max", "tokens_per_sec", f.s["fig11_title_b"], 1, "max"))
    for ax, rel_col, abs_col, title, nd, best in panels:
        rel = 100 * ts.pivot(index="model_key", columns="n_threads", values=rel_col).loc[models, threads]
        absolute = ts.pivot(index="model_key", columns="n_threads", values=abs_col).loc[models, threads]
        vmin, vmax = rel.values.min(), rel.values.max()
        ax.imshow(rel.values, cmap=cmap, aspect="auto", vmin=vmin - 0.15 * (vmax - vmin), vmax=vmax, zorder=1)
        for i, k in enumerate(models):
            best_j = int(np.argmin(rel.loc[k].values) if best == "min" else np.argmax(rel.loc[k].values))
            for j, n in enumerate(threads):
                shade = (rel.iloc[i, j] - vmin) / (vmax - vmin) if vmax > vmin else 0
                ax.text(j, i, f"{f.num(rel.iloc[i, j], 0)} %\n({f.num(absolute.iloc[i, j], nd)})", ha="center",
                        va="center", fontsize=8.5, color="white" if shade > 0.62 else INK,
                        fontweight="bold" if j == best_j else "normal", zorder=3)
            ax.add_patch(Rectangle((best_j - 0.5, i - 0.5), 1, 1, fill=False, edgecolor=INK, linewidth=2, zorder=4))
        ax.set_xticks(range(len(threads)))
        ax.set_xticklabels([str(n) for n in threads])
        ax.set_yticks(range(len(models)))
        ax.set_yticklabels(models, fontsize=8.5)
        ax.set_xlabel(f.s["threads"], color=MUTED)
        ax.set_title(title, color=INK, fontsize=11.5, pad=10)
        ax.tick_params(colors=MUTED, length=0)
        for side in ax.spines.values():
            side.set_visible(False)
        # White gaps between cells (2 px spacer)
        for j in range(1, len(threads)):
            ax.axvline(j - 0.5, color="white", linewidth=2, zorder=2)
        for i in range(1, len(models)):
            ax.axhline(i - 0.5, color="white", linewidth=2, zorder=2)
    fig.text(0.5, 0.005, f.s["fig11_note"], ha="center", va="bottom", color=MUTED, fontsize=7.5)
    fig.tight_layout(rect=(0, 0.1, 1, 1), w_pad=4)
    return save(fig, out_dir, f.s["fig_names"][11])


def table_T9(data, f):
    t = data.night_turbo.set_index(["model_key", "turbo"])
    r = data.night_turbo_ratio
    tr = f.lang == "tr"
    rows = []
    for k in [k for k in MODEL_ORDER if k in r.index]:
        off, on = t.loc[(k, "off")], t.loc[(k, "on")]
        rows.append({
            "Model": k,
            ("J/token kapalı" if tr else "J/token off"): f.pm(off["J_per_token"], off["J_per_token_sd"], 3),
            ("J/token açık" if tr else "J/token on"): f.pm(on["J_per_token"], on["J_per_token_sd"], 3),
            ("Oran J [%95 GA]" if tr else "Ratio J [95 % CI]"): f"{f.num(r.loc[k, 'J_ratio_on_off'], 3)} {f.ci(r.loc[k, 'J_ratio_ci95_low'], r.loc[k, 'J_ratio_ci95_high'], 3)}",
            ("Token/s kapalı" if tr else "Tokens/s off"): f.pm(off["tokens_per_sec"], off["tokens_per_sec_sd"], 2),
            ("Token/s açık" if tr else "Tokens/s on"): f.pm(on["tokens_per_sec"], on["tokens_per_sec_sd"], 2),
            ("Oran hız [%95 GA]" if tr else "Ratio speed [95 % CI]"): f"{f.num(r.loc[k, 'speed_ratio_on_off'], 3)} {f.ci(r.loc[k, 'speed_ratio_ci95_low'], r.loc[k, 'speed_ratio_ci95_high'], 3)}",
            ("Güç kapalı (W)" if tr else "Power off (W)"): f.pm(off["cpu_power_W"], off["cpu_power_W_sd"], 1),
            ("Güç açık (W)" if tr else "Power on (W)"): f.pm(on["cpu_power_W"], on["cpu_power_W_sd"], 1),
            ("Oran güç [%95 GA]" if tr else "Ratio power [95 % CI]"): f"{f.num(r.loc[k, 'power_ratio_on_off'], 3)} {f.ci(r.loc[k, 'power_ratio_ci95_low'], r.loc[k, 'power_ratio_ci95_high'], 3)}",
            ("Frekans kapalı / açık (MHz)" if tr else "Frequency off / on (MHz)"): f"{f.num(off['freq_mean_mhz'], 0)} / {f.num(on['freq_mean_mhz'], 0)}",
            ("Tepe sıcaklık kapalı / açık (°C)" if tr else "Peak temperature off / on (°C)"): f"{f.num(off['temp_max_C'], 0)} / {f.num(on['temp_max_C'], 0)}",
            ("mgCO2eq/token kapalı / açık" if tr else "mgCO2eq/token off / on"): f"{f.num(off['mgco2eq_per_token'], 2)} / {f.num(on['mgco2eq_per_token'], 2)}",
        })
    return pd.DataFrame(rows)


def table_T10(data, f):
    ts = data.night_threads
    tr = f.lang == "tr"
    rows = []
    for k in [k for k in MODEL_ORDER if k in set(ts["model_key"])]:
        g = ts[ts["model_key"] == k].set_index("n_threads").sort_index()
        j, s, u = g["J_per_token"], g["tokens_per_sec"], g["utilisation_est2_pct"]
        row = {"Model": k, ("Parametre (B)" if tr else "Params (B)"): f.num(MODELS[k]["params_b"], 1)}
        for n in g.index:
            row[f"J/token @{n}"] = f.pm(j[n], g.loc[n, "J_per_token_sd"], 3)
        for n in g.index:
            row[("Token/s" if tr else "Tokens/s") + f" @{n}"] = f.pm(s[n], g.loc[n, "tokens_per_sec_sd"], 2)
        row.update({
            ("En düşük enerji" if tr else "Min energy"): int(j.idxmin()),
            ("Doyum (%98)" if tr else "Saturation (98 %)"): int(s[s >= 0.98 * s.max()].index.min()),
            ("En yüksek hız" if tr else "Max speed"): int(s.idxmax()),
            ("Hız 2→4 (%)" if tr else "Speed 2→4 (%)"): f.num(100 * (s[4] / s[2] - 1), 1),
            ("Hız 4→8 (%)" if tr else "Speed 4→8 (%)"): f.num(100 * (s[8] / s[4] - 1), 1),
            ("Enerji 2→8 (%)" if tr else "Energy 2→8 (%)"): f.num(100 * (j[8] / j[2] - 1), 1),
            ("Bant genişliği kullanımı 2 / 4 / 8 (%)" if tr else "Bandwidth utilisation 2 / 4 / 8 (%)"):
                " / ".join(f.num(u[n], 0) for n in g.index),
        })
        rows.append(row)
    return pd.DataFrame(rows)


NIGHT_TABLE_NOTES = {
    "tr": {
        1: "Perplexity ve bit/bayt: WikiText-2 testinin ilk 90 parçası (gece koşusu night_01, 23.040 puanlanan token); "
           "oturum 01'deki 60 parçalık alt küme sonuçları docs/findings_report.md Ek A'dadır.",
        "1_full": "Perplexity ve bit/bayt: WikiText-2 testinin tamamı (564 parça, 144.384 puanlanan token; ilk 90 parça "
                  "night_01, kalanı night_02, parça düzeyinde birleştirilmiş). 60 ve 90 parçalık alt küme sonuçları "
                  "docs/findings_report.md Ek A'dadır.",
        "10_2s": "İki bağımsız oturum (night_01, night_02), 8 iş parçacığı, oturum başına 10 tur; Turbo açık/kapalı birimler "
                 "aynı turda iç içe. Değerler: oturum ortalamalarının ortalaması ± oturumlar arası SS (n = 2). Oranlar: 20 "
                 "turda eşleştirilmiş açık/kapalı, GA Student-t (df = 19). Kapalı = azami işlemci durumu %99, açık = %100.",
        2: NIGHT_MARK + " Gece koşusu (night_01): tek oturum, 8 iş parçacığı, 10 tur; oran her turda eşleştirilmiş, GA "
           "turlar üzerinden Student-t (df = 9). Aynı gecede Q8_0 / Q4_K_M oranı kampanya oranıyla (1,600) örtüşür.",
        8: NIGHT_MARK + " Gece koşusunda ölçüldü (night_01, 10 tur); token/s çapa modeller (8B-Q4_K_M, 8B-Q8_0) üzerinden "
           "kampanya ölçeğine getirildi; SS turlar arasıdır.",
        10: "Tek oturum (night_01), 8 iş parçacığı, 10 tur; Turbo açık/kapalı birimler aynı turda iç içe. Oranlar turda "
            "eşleştirilmiş açık/kapalı; GA Student-t (df = 9). Kapalı = azami işlemci durumu %99, açık = %100.",
        11: "Tek oturum (night_01), 5 tur × 3 istem, 100 token/üretim. Doyum: en yüksek hızın %98'ine ulaşan en küçük iş "
            "parçacığı sayısı. Bant genişliği kullanımı: düzeltilmiş bayt/token × token/s / 34,1 GB/s (bkz. T7).",
    },
    "en": {
        1: "Perplexity and bits/byte: the first 90 chunks of the WikiText-2 test set (night run night_01, 23,040 scored "
           "tokens); the 60-chunk subset results of session 01 are in docs/findings_report.md, Appendix A.",
        "1_full": "Perplexity and bits/byte: the whole WikiText-2 test set (564 chunks, 144,384 scored tokens; the first 90 "
                  "chunks from night_01, the rest from night_02, merged chunk by chunk). The 60- and 90-chunk subset results "
                  "are in docs/findings_report.md, Appendix A.",
        "10_2s": "Two independent sessions (night_01, night_02), 8 threads, 10 rounds per session; Turbo-on/off units "
                 "interleaved within each round. Values: mean of the session means ± between-session SD (n = 2). Ratios: "
                 "paired on/off over 20 rounds, CI: Student-t (df = 19). Off = maximum processor state 99 %, on = 100 %.",
        2: NIGHT_MARK + " Night run (night_01): one session, 8 threads, 10 rounds; the ratio is paired within each round, the "
           "CI is a Student-t interval over rounds (df = 9). In the same night the Q8_0 / Q4_K_M ratio matches the campaign (1.600).",
        8: NIGHT_MARK + " Measured in the night run (night_01, 10 rounds); tokens/s scaled to the campaign through the anchor "
           "models (8B-Q4_K_M, 8B-Q8_0); the SD is between rounds.",
        10: "One session (night_01), 8 threads, 10 rounds; Turbo-on/off units interleaved within each round. Ratios are "
            "paired on/off within a round; CI: Student-t (df = 9). Off = maximum processor state 99 %, on = 100 %.",
        11: "One session (night_01), 5 rounds × 3 prompts, 100 tokens per generation. Saturation: the smallest thread count "
            "reaching 98 % of the maximum speed. Bandwidth utilisation: corrected bytes/token × tokens/s / 34.1 GB/s (see T7).",
    },
}


def table_T8(data, f):
    """Prefill vs decode: speed and CPU energy per token (prefill per prompt token)."""
    e = data.energy_table()
    tr = f.lang == "tr"
    prompt_tokens = data.energy.groupby("model_key")["prompt_tokens"].agg(["min", "max"])
    rows = []
    for k in [k for k in MODEL_ORDER if k in e["tokens_per_sec"].index]:
        pf_j, dc_j = e["prefill_J_per_prompt_token"].loc[k], e["cpu_J_per_token"].loc[k]
        rows.append({
            "Model": k + (" (base)" if MODELS[k]["base"] else ""),
            ("İstem (token)" if tr else "Prompt (tokens)"): f"{int(prompt_tokens.loc[k, 'min'])}–{int(prompt_tokens.loc[k, 'max'])}",
            ("Prefill token/s" if tr else "Prefill tokens/s"): f.pm(e["prefill_tokens_per_sec"].loc[k, "mean"], e["prefill_tokens_per_sec"].loc[k, "sd"], 0),
            ("Decode token/s" if tr else "Decode tokens/s"): f.pm(e["tokens_per_sec"].loc[k, "mean"], e["tokens_per_sec"].loc[k, "sd"], 2),
            ("Prefill J / istem token'ı" if tr else "Prefill J / prompt token"): f.pm(pf_j["mean"], pf_j["sd"], 3),
            ("Decode J / token" if tr else "Decode J / token"): f.pm(dc_j["mean"], dc_j["sd"], 3),
            ("Decode / prefill (J)" if tr else "Decode / prefill (J)"): f.num(dc_j["mean"] / pf_j["mean"], 1),
        })
    return pd.DataFrame(rows)


# ---------------------------------------------------------------- main
def run(lang, results_dir, fig_root, table_root):
    f = Fmt(lang)
    data = Data(results_dir)
    fig_dir, table_dir = os.path.join(fig_root, lang), os.path.join(table_root, lang)
    produced = []
    figures = [fig01_energy_speed, fig02_tradeoff, fig03_thread_scaling, fig04_thread_pareto,
               fig05_session_consistency, fig06_prefill_decode, fig07_bandwidth_models, fig08_bandwidth_threads,
               fig09_roofline]
    if data.night is not None:
        figures += [fig10_turbo, fig11_thread_optimum]
    for fn in figures:
        produced += fn(data, f, fig_dir)
    titles_tr = {1: "Tablo 1. Ana sonuçlar", 2: "Tablo 2. Nicemleme karşılaştırması", 3: "Tablo 3. Seyrek – yoğun eşdeğerlik",
                 4: "Tablo 4. İş parçacığı ölçekleme", 5: "Tablo 5. Oturumlar arası varyans ayrıştırması",
                 6: "Tablo 6. Deney koşulları", 7: "Tablo 6b. Oturumlar", 8: "Tablo 7. Bellek bant genişliği kullanımı",
                 9: "Tablo 8. Prefill ve decode: hız ve enerji",
                 10: "Tablo 9. Turbo Boost açık / kapalı (kontrollü deney, night_01)",
                 11: "Tablo 10. Enerji-optimal iş parçacığı sayısı, 8 model (night_01)"}
    titles_en = {1: "Table 1. Main results", 2: "Table 2. Quantization comparison", 3: "Table 3. Sparse vs dense equivalence",
                 4: "Table 4. Thread scaling", 5: "Table 5. Between-session variance decomposition",
                 6: "Table 6. Experimental conditions", 7: "Table 6b. Sessions", 8: "Table 7. Memory-bandwidth utilisation",
                 9: "Table 8. Prefill vs decode: speed and energy",
                 10: "Table 9. Turbo Boost on / off (controlled experiment, night_01)",
                 11: "Table 10. Energy-optimal thread count, 8 models (night_01)"}
    titles = titles_tr if lang == "tr" else titles_en
    if data.turbo_sessions >= 2:
        titles[10] = titles[10].replace("night_01", "night_01 + night_02")
    builders = [(1, table_T1), (2, table_T2), (3, table_T3), (4, table_T4), (5, table_T5), (6, table_T6),
                (7, table_T6b), (8, table_T7), (9, table_T8)]
    if data.night is not None:
        builders += [(10, table_T9), (11, table_T10)]
    for i, builder in builders:
        note = T6_NOTE[lang] if builder is table_T6 else None
        if data.night is not None and i in NIGHT_TABLE_NOTES[lang]:
            note = NIGHT_TABLE_NOTES[lang][i]
            if i == 1 and getattr(data, "ppl_chunks", 0) >= 564:
                note = NIGHT_TABLE_NOTES[lang]["1_full"]
            if i == 10 and data.turbo_sessions >= 2:
                note = NIGHT_TABLE_NOTES[lang]["10_2s"]
        write_table(builder(data, f), table_dir, f.s["table_names"][i], f, titles[i], note=note)
        produced.append(os.path.join(table_dir, f.s["table_names"][i] + ".{csv,md,docx}"))
    return produced


def main():
    parser = argparse.ArgumentParser(description="Figures and tables of the Phase 2 campaign (tr / en).")
    parser.add_argument("--lang", choices=["tr", "en", "all"], default="tr")
    parser.add_argument("--results-dir", default=RESULTS_DIR)
    parser.add_argument("--figures-dir", default=os.path.join(PROJECT_ROOT, "figures"))
    parser.add_argument("--tables-dir", default=os.path.join(PROJECT_ROOT, "tables"))
    args = parser.parse_args()
    plt.rcParams.update({"font.size": 10, "font.family": "DejaVu Sans", "axes.titlecolor": INK, "hatch.linewidth": 0.6})
    for lang in (["tr", "en"] if args.lang == "all" else [args.lang]):
        for path in run(lang, args.results_dir, args.figures_dir, args.tables_dir):
            print("written:", path)


if __name__ == "__main__":
    main()
